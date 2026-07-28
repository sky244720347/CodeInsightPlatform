#!/usr/bin/env python3
"""
Excel A/B 交叉标记（约 60 万行可用）

功能：
1. 用 A 的数据列去查 B 的两个数据列；命中任一列 → A 标记列写「是」，否则「否」
2. 用 B 的两个数据列去查 A 的数据列；该行任一列命中 → B 标记列写「是」，否则「否」

依赖：pip install openpyxl
用法示例：
  python excel_cross_mark.py ^
    --a A.xlsx --b B.xlsx ^
    --a-data 数据 --a-mark 标记 ^
    --b-data1 数据1 --b-data2 数据2 --b-mark 标记 ^
    --out-a A_marked.xlsx --out-b B_marked.xlsx

也可用列索引（0 起）：--a-data 0 --a-mark 1 --b-data1 0 --b-data2 1 --b-mark 2
"""

from __future__ import annotations

import argparse
import sys
import time
from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.worksheet import Worksheet


YES = "是"
NO = "否"


def _norm(v: Any) -> str | None:
    """统一成可比较的字符串；空值返回 None（不参与命中）。"""
    if v is None:
        return None
    if isinstance(v, str):
        s = v.strip()
        return s if s else None
    # Excel 数字可能是 int/float；去掉 .0 避免 123 vs 123.0 对不上
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip() or None


def _resolve_col(header_row: tuple[Any, ...], spec: str) -> int:
    """列名或 0-based 列索引 → 0-based 列下标。"""
    if spec.isdigit():
        idx = int(spec)
        if idx < 0 or idx >= len(header_row):
            raise SystemExit(f"列索引越界: {spec}（表头共 {len(header_row)} 列）")
        return idx
    for i, cell in enumerate(header_row):
        if cell is not None and str(cell).strip() == spec:
            return i
    raise SystemExit(f"找不到列名「{spec}」。表头={list(header_row)}")


def _iter_data_rows(ws: Worksheet) -> Iterable[tuple[Any, ...]]:
    rows = ws.iter_rows(values_only=True)
    header = next(rows, None)
    if header is None:
        raise SystemExit("工作表为空")
    yield header  # type: ignore[misc]
    for row in rows:
        yield row


def load_a(
    path: Path,
    sheet: str | None,
    data_spec: str,
    mark_spec: str,
) -> tuple[list[tuple[Any, ...]], int, int, set[str]]:
    """读 A：返回全部行(含表头)、数据列下标、标记列下标、数据集合。"""
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet] if sheet else wb.active
    assert ws is not None

    it = _iter_data_rows(ws)
    header = next(it)
    data_i = _resolve_col(header, data_spec)
    mark_i = _resolve_col(header, mark_spec)

    rows: list[tuple[Any, ...]] = [header]
    a_set: set[str] = set()
    for row in it:
        rows.append(row)
        key = _norm(row[data_i] if data_i < len(row) else None)
        if key is not None:
            a_set.add(key)
    wb.close()
    return rows, data_i, mark_i, a_set


def load_b(
    path: Path,
    sheet: str | None,
    data1_spec: str,
    data2_spec: str,
    mark_spec: str,
) -> tuple[list[tuple[Any, ...]], int, int, int, set[str]]:
    """读 B：返回全部行、两数据列下标、标记列下标、两列并集。"""
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet] if sheet else wb.active
    assert ws is not None

    it = _iter_data_rows(ws)
    header = next(it)
    d1 = _resolve_col(header, data1_spec)
    d2 = _resolve_col(header, data2_spec)
    mark_i = _resolve_col(header, mark_spec)

    rows: list[tuple[Any, ...]] = [header]
    b_set: set[str] = set()
    for row in it:
        rows.append(row)
        for idx in (d1, d2):
            key = _norm(row[idx] if idx < len(row) else None)
            if key is not None:
                b_set.add(key)
    wb.close()
    return rows, d1, d2, mark_i, b_set


def _ensure_width(row: tuple[Any, ...], min_len: int) -> list[Any]:
    cells = list(row)
    if len(cells) < min_len:
        cells.extend([None] * (min_len - len(cells)))
    return cells


def mark_a(
    rows: list[tuple[Any, ...]],
    data_i: int,
    mark_i: int,
    b_set: set[str],
) -> list[list[Any]]:
    out: list[list[Any]] = []
    header = _ensure_width(rows[0], mark_i + 1)
    out.append(header)
    need = max(data_i, mark_i) + 1
    for row in rows[1:]:
        cells = _ensure_width(row, need)
        key = _norm(cells[data_i])
        cells[mark_i] = YES if (key is not None and key in b_set) else NO
        out.append(cells)
    return out


def mark_b(
    rows: list[tuple[Any, ...]],
    d1: int,
    d2: int,
    mark_i: int,
    a_set: set[str],
) -> list[list[Any]]:
    out: list[list[Any]] = []
    header = _ensure_width(rows[0], mark_i + 1)
    out.append(header)
    need = max(d1, d2, mark_i) + 1
    for row in rows[1:]:
        cells = _ensure_width(row, need)
        k1 = _norm(cells[d1])
        k2 = _norm(cells[d2])
        hit = (k1 is not None and k1 in a_set) or (k2 is not None and k2 in a_set)
        cells[mark_i] = YES if hit else NO
        out.append(cells)
    return out


def write_xlsx(path: Path, rows: list[list[Any]]) -> None:
    """write_only 流式写出，适合大表。"""
    wb = Workbook(write_only=True)
    ws = wb.create_sheet()
    for row in rows:
        ws.append(row)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def main() -> int:
    p = argparse.ArgumentParser(description="Excel A/B 交叉标记（集合查找）")
    p.add_argument("--a", required=True, help="A 表路径")
    p.add_argument("--b", required=True, help="B 表路径")
    p.add_argument("--a-sheet", default=None, help="A 工作表名（默认第一个）")
    p.add_argument("--b-sheet", default=None, help="B 工作表名（默认第一个）")
    p.add_argument("--a-data", required=True, help="A 数据列名或索引")
    p.add_argument("--a-mark", required=True, help="A 标记列名或索引")
    p.add_argument("--b-data1", required=True, help="B 数据列1 名或索引")
    p.add_argument("--b-data2", required=True, help="B 数据列2 名或索引")
    p.add_argument("--b-mark", required=True, help="B 标记列名或索引")
    p.add_argument("--out-a", default="A_marked.xlsx", help="输出 A（默认 A_marked.xlsx）")
    p.add_argument("--out-b", default="B_marked.xlsx", help="输出 B（默认 B_marked.xlsx）")
    args = p.parse_args()

    a_path, b_path = Path(args.a), Path(args.b)
    if not a_path.exists():
        print(f"找不到 A: {a_path}", file=sys.stderr)
        return 1
    if not b_path.exists():
        print(f"找不到 B: {b_path}", file=sys.stderr)
        return 1

    t0 = time.perf_counter()
    print("读取 A ...")
    a_rows, a_data_i, a_mark_i, a_set = load_a(
        a_path, args.a_sheet, args.a_data, args.a_mark
    )
    print(f"  A 行数(含表头)={len(a_rows)}, 去重键={len(a_set)}")

    print("读取 B ...")
    b_rows, b_d1, b_d2, b_mark_i, b_set = load_b(
        b_path, args.b_sheet, args.b_data1, args.b_data2, args.b_mark
    )
    print(f"  B 行数(含表头)={len(b_rows)}, 两列并集键={len(b_set)}")

    print("标记 A ...")
    marked_a = mark_a(a_rows, a_data_i, a_mark_i, b_set)
    print("标记 B ...")
    marked_b = mark_b(b_rows, b_d1, b_d2, b_mark_i, a_set)

    out_a, out_b = Path(args.out_a), Path(args.out_b)
    print(f"写出 {out_a} ...")
    write_xlsx(out_a, marked_a)
    print(f"写出 {out_b} ...")
    write_xlsx(out_b, marked_b)

    elapsed = time.perf_counter() - t0
    print(f"完成，耗时 {elapsed:.1f}s")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
