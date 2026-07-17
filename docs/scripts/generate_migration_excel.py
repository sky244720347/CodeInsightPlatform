"""Generate CodeInsight migration plan Excel — module + page division (owners blank)."""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OUTPUT = r"c:\project\codeInsight\CodeInsightPlatform\docs\dual-team-migration-plan.xlsx"

HEADER_FILL = PatternFill("solid", fgColor="4472C4")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
SECTION_FILL = PatternFill("solid", fgColor="D9E1F2")
SECTION_FONT = Font(bold=True, size=12)
MODULE_FILL = PatternFill("solid", fgColor="E2EFDA")
PAGE_FILL = PatternFill("solid", fgColor="DDEBF7")
IMPORTANT_FILL = PatternFill("solid", fgColor="FCE4D6")
THIN = Side(style="thin", color="B4B4B4")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")

HEADERS = [
    "分工类型",
    "模块",
    "菜单分组",
    "父菜单",
    "页面/能力",
    "迁移内容",
    "备注",
    "负责人",
]
COL_WIDTHS = [10, 16, 14, 16, 18, 44, 28, 10]

# (模块, 工作项, 迁移内容, 备注, 是否重要)
MODULE_ROWS = [
    ("auth", "PACAS 登录对接", "接入平安科技身份认证；Session/操作人上下文；替换现有登录", "系归组对接", True),
    ("auth", "UAMP 权限对接", "菜单/按钮/API 级权限拦截与查询", "系归组对接", False),
    ("task", "任务状态机与调度", "创建/列表/详情/队列/取消/进度；全状态机流转", "", False),
    ("task", "神兵任务入队 API", "供系归组 Webhook 回调创建增量/扫描任务", "跨组接口", True),
    ("scanner", "代码拉取与扫描", "Git 拉取、增量 diff、文件 snapshot、扫描试跑", "依赖 Git 凭据服务", False),
    ("parser", "Java 静态解析", "JavaParser AST 解析、类/方法/注解提取", "", False),
    ("callchain", "调用链持久化", "ci_method_call 落表；增量删重建", "", False),
    ("entrypoint", "入口识别与复核", "入口发现落表；复核确认/驳回 API", "", False),
    ("hierarchy", "模块层级与复核", "模块树构建；层级复核保存", "", False),
    ("ai", "AI 归纳调用", "模块层级/功能文档生成；Token 审计落表", "", False),
    ("draft", "草稿工作区", "Markdown 草稿读写；编辑锁；自动保存；确认", "统一缓存", False),
    ("knowledge", "知识查看与纠错", "入口/层级/文档三页；纠错重跑；NAS 修订", "", False),
    ("push", "知识推送与版本", "NAS/Git 发布；版本快照；ZIP 导出", "", False),
    ("businessknowledge", "业务知识维护", "注入 AI 提示词的业务知识 CRUD", "", False),
    ("system", "系统管理", "业务系统 CRUD；负责人；扫描范围配置", "", False),
    ("repository", "仓库管理", "Git 仓库绑定；扫描规则；lastCommit 基线", "", False),
    ("prompt", "提示词库", "提示词模板 CRUD；绑定；试跑", "", False),
    ("model", "模型配置", "AI 模型注册；连通性测试；公司配额联动", "", False),
    ("dashboard", "运营看板", "任务/流水线/覆盖率/模型用量统计 API", "", False),
    ("token", "Token 审计", "AI 调用明细与费用统计", "", False),
    ("log", "操作日志", "ci_operation_log 检索与展示", "", False),
    ("quotacontrol", "流量控制", "全局/系统级配额与并发限制", "", False),
    ("scanwindow", "扫描编排", "扫描时间窗口配置；热力图数据", "触发依赖神兵", False),
    ("common", "公共基础设施", "异常/响应/配置刷新；公司框架 Parent 适配", "", False),
    ("storage", "存储层", "本地/NAS 路径；草稿与知识正文读写", "", False),
    ("cluster", "集群与分布式", "Leader 选举；任务认领；统一缓存替代 Redis 锁", "统一缓存", False),
    ("integration", "Git 代码库权限", "公司 Git 权限 API；JGit 凭据注入", "系归组", True),
    ("integration", "神兵 Webhook", "接收发布状态；幂等；回调任务入队", "系归组", True),
    ("integration", "CLI 知识聚合 API", "HTTP 读 NAS 供命令行批量下载知识", "系归组", False),
    ("integration", "知识打通 ai_workspace", "定时同步知识至外部平台", "联系人：杨东", False),
    ("infra", "Apollo 全量配置", "全部 namespace 申请、录入、TEST/PRO 发布", "征信组维护", True),
    ("infra", "统一缓存对接", "草稿锁/集群锁/分布式信号量 SDK 接入", "", True),
    ("infra", "资源申请", "PostgreSQL / 统一缓存 / Apollo / NAS 测试与生产", "", False),
    ("infra", "Bettle + DB 初始化", "32 表脚本；增量 DDL 走 Bettle 审批", "", False),
    ("infra", "JAR 私服适配", "公司 BOM；JGit/JavaParser/apollo-client 白名单", "", False),
    ("pipeline", "后端部署流水线", "JAR/镜像 CI/CD；测试/生产 Job", "", True),
    ("pipeline", "前端部署流水线", "Vite build；环境变量；静态资源；网关联调", "", True),
    ("frontend", "Admin 壳与全局框架", "公司 Admin 布局；Hash/Browser 路由；顶栏/侧栏", "", False),
]

# (模块, 菜单分组, 父菜单, 页面, 迁移内容, 备注)
PAGE_ROWS = [
    ("auth", "—", "—", "登录页", "PACAS 登录页/回调；Token 存储；跳转首页", "/login"),
    ("frontend", "—", "—", "整体 Layout", "公司 Admin 布局；UAMP 控制可见菜单；退出登录", "全站框架"),
    ("pipeline", "—", "—", "前端流水线", "Vite CI；VITE_API_BASE_URL 注入；CDN/网关部署", "非业务页"),
    ("dashboard", "仪表盘 / 看板", "—", "任务概览", "ECharts 图表；对接 dashboard 任务统计 API", "/dashboard/tasks"),
    ("dashboard", "仪表盘 / 看板", "—", "AI 模型用量", "图表+筛选；dashboard + token 数据", "/dashboard/ai-usage"),
    ("dashboard", "仪表盘 / 看板", "—", "流水线分析", "阶段耗时柱状图；dashboard + 任务日志", "/dashboard/pipeline"),
    ("dashboard", "仪表盘 / 看板", "—", "系统覆盖报表", "覆盖率表格/图；dashboard 覆盖率 API", "/dashboard/coverage"),
    ("token", "仪表盘 / 看板", "—", "Token 审计", "分页表格+图表；token 模块 API", "/audit"),
    ("log", "仪表盘 / 看板", "—", "操作日志", "日志检索页；log 模块 API", "/logs"),
    ("system", "知识管理", "—", "系统与仓库", "系统/仓库 CRUD；扫描配置；试跑抽屉", "/systems"),
    ("repository", "知识管理", "—", "系统与仓库", "同上（仓库子功能）", "与 system 同页"),
    ("task", "知识管理", "知识任务构建", "任务列表", "任务表格+筛选；task 列表/状态 API", "/tasks/query"),
    ("task", "知识管理", "知识任务构建", "任务详情", "详情页+日志轮询+进度 Steps", "/tasks/:id"),
    ("task", "知识管理", "知识任务构建", "手动下发任务", "多步骤创建表单；task 创建 API", "/tasks/dispatch"),
    ("task", "知识管理", "任务管控", "队列查看", "队列视图；PENDING 任务调度 API", "/tasks/queue"),
    ("entrypoint", "知识管理", "知识任务构建", "扫描入口复核", "入口清单树/列表；确认或驳回", "/tasks/entrypoint-review"),
    ("hierarchy", "知识管理", "知识任务构建", "知识模块复核", "模块树编辑器；层级复核保存", "/tasks/hierarchy-review"),
    ("draft", "知识管理", "知识任务构建", "生成知识复核", "Monaco+模块目录树；草稿锁/保存/确认", "/drafts"),
    ("push", "知识管理", "知识任务构建", "知识推送", "版本列表+推送确认+ZIP 导出", "/push"),
    ("knowledge", "知识管理", "知识查询", "扫描入口", "入口列表+纠错模式；已发布入口查询", "/knowledge/entrypoints"),
    ("knowledge", "知识管理", "知识查询", "模块层级", "层级树+纠错抽屉；已发布层级查询", "/knowledge/hierarchy"),
    ("knowledge", "知识管理", "知识查询", "知识文档", "树形/列表+Markdown 预览+编辑+纠错", "/knowledge/documents"),
    ("model", "基础配置", "—", "模型配置", "模型 CRUD+试跑；model 模块", "/basic/models"),
    ("prompt", "基础配置", "—", "提示词", "提示词编辑+试跑；prompt 模块", "/basic/prompts"),
    ("auth", "基础配置", "—", "权限管理", "UAMP 权限展示；权限查询 API", "/basic/permissions"),
    ("quotacontrol", "基础配置", "任务管控", "流量配置", "配额表单；quotacontrol 模块", "/basic/quota"),
    ("scanwindow", "基础配置", "任务管控", "任务编排", "热力图+cron UI；scanwindow 配置", "/basic/orchestration"),
]


def style_header_row(ws, row, ncol):
    for c in range(1, ncol + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER


def write_rows(ws, start_row, rows):
    ncol = len(HEADERS)
    r = start_row
    for row_data in rows:
        important = row_data.pop("_important", False)
        row_type = row_data[0]
        for i, val in enumerate(row_data, 1):
            cell = ws.cell(row=r, column=i, value=val)
            cell.alignment = WRAP
            cell.border = BORDER
        fill = IMPORTANT_FILL if important else (MODULE_FILL if row_type == "模块" else PAGE_FILL)
        for c in range(1, ncol + 1):
            ws.cell(row=r, column=c).fill = fill
        r += 1
    return r


def build_overview_sheet(wb):
    ws = wb.active
    ws.title = "分工总览"

    ws["A1"] = "CodeInsight 公司框架迁移 — 分工总览（模块 + 页面）"
    ws.merge_cells("A1:H1")
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = (
        "绿色=模块分工 · 蓝色=页面分工 · 负责人列留空待填写 · "
        "菜单结构对齐 frontend BasicLayout 侧栏"
    )
    ws.merge_cells("A2:H2")
    ws["A2"].font = Font(size=10, color="666666")

    row = 4
    for i, h in enumerate(HEADERS, 1):
        ws.cell(row=row, column=i, value=h)
    style_header_row(ws, row, len(HEADERS))
    row += 1

    combined = []
    for mod, name, content, note, important in MODULE_ROWS:
        combined.append(
            {
                "_important": important,
                "data": ["模块", mod, "", "", name, content, note, ""],
            }
        )
    for mod, group, parent, page, content, note in PAGE_ROWS:
        combined.append(
            {
                "_important": False,
                "data": ["页面", mod, group, parent, page, content, note, ""],
            }
        )

    for item in combined:
        important = item["_important"]
        row_data = item["data"]
        for i, val in enumerate(row_data, 1):
            cell = ws.cell(row=row, column=i, value=val)
            cell.alignment = WRAP
            cell.border = BORDER
        fill = IMPORTANT_FILL if important else (MODULE_FILL if row_data[0] == "模块" else PAGE_FILL)
        for c in range(1, len(HEADERS) + 1):
            ws.cell(row=row, column=c).fill = fill
        row += 1

    for i, w in enumerate(COL_WIDTHS, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A5"


def build_schedule_sheet(wb):
    import sys
    from pathlib import Path

    sys.path.insert(0, str(Path(__file__).resolve().parent))
    from update_migration_schedule import build_schedule_sheet as fill_schedule

    ws = wb.create_sheet("排期")
    fill_schedule(ws)


def main():
    wb = Workbook()
    build_overview_sheet(wb)
    build_schedule_sheet(wb)
    wb.save(OUTPUT)
    print(f"Saved: {OUTPUT}")


if __name__ == "__main__":
    main()
