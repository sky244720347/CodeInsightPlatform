"""Update only the 排期 sheet in migration Excel; preserve user-edited 分工总览."""
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

XLSX = r"c:\project\codeInsight\CodeInsightPlatform\docs\dual-team-migration-plan.xlsx"
XLSX_CN = r"c:\project\codeInsight\CodeInsightPlatform\docs\CodeInsight迁移方案-分工与排期.xlsx"

HEADER_FILL = PatternFill("solid", fgColor="4472C4")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
SECTION_FILL = PatternFill("solid", fgColor="D9E1F2")
SECTION_FONT = Font(bold=True, size=12)
THIN = Side(style="thin", color="B4B4B4")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")

HEADERS = ["天数", "模块", "工作项", "交付内容", "依赖/备注", "完成标志", "负责人"]
COL_WIDTHS = [6, 14, 22, 40, 26, 20, 10]

# W1：测试环境部署成功 + 全量扫描主链路打通（INITIAL 任务 → 至少 1 篇草稿）
W1_ROWS = [
    ["D1", "infra", "资源申请 / JAR 适配", "PG、统一缓存、Apollo、NAS 测试资源提单；公司 BOM 与白名单", "", "申请单提交", ""],
    ["D1", "infra", "Bettle + DB 初始化", "32 表 DDL 脚本整理并提交 Bettle（测试环境）", "", "脚本已提交", ""],
    ["D1", "pipeline", "后端部署流水线", "测试环境 CI/CD Job、镜像构建配置", "", "流水线可触发", ""],
    ["D1", "pipeline", "前端部署流水线", "Vite build Job、静态资源上传骨架", "", "流水线可触发", ""],
    ["D1", "integration", "Git 代码库权限", "权限 API 调研；单测试仓库凭据方案", "系归组", "", ""],
    ["D1", "auth", "PACAS 登录对接", "测试环境 PACAS 应用注册与回调地址申请", "系归组", "文档就绪", ""],
    ["D2", "infra", "Apollo 全量配置", "Apollo TEST 接入；datasource / storage / ai 等核心 namespace", "", "配置可读取", ""],
    ["D2", "infra", "统一缓存对接", "统一缓存 SDK 接入；连通性验证", "", "缓存 ping 通", ""],
    ["D2", "infra", "Bettle + DB 初始化", "测试库执行 DDL；基础数据可连", "", "库表就绪", ""],
    ["D2", "common / storage", "公共基础设施 / 存储层", "common 配置；NAS/本地 storage 路径验收", "", "", ""],
    ["D2", "auth", "PACAS 登录对接", "测试 SSO 登录打通；Session 写入", "", "可登录", ""],
    ["D2", "frontend", "登录页 / 整体 Layout", "PACAS 回调页；Admin 壳骨架", "", "登录跳转首页", ""],
    ["D3", "pipeline", "后端部署流水线", "后端 JAR 部署测试环境首跑", "infra 就绪", "后端服务可访问", ""],
    ["D3", "system", "系统管理", "业务系统 CRUD", "", "可登记系统", ""],
    ["D3", "repository", "仓库管理", "Git 仓库绑定；扫描范围配置", "integration Git 凭据", "可绑定仓库", ""],
    ["D3", "integration", "Git 代码库权限", "单测试仓库凭据注入 JGit", "", "拉代码成功", ""],
    ["D3", "prompt", "提示词库", "默认提示词配置（全量任务可用）", "", "", ""],
    ["D3", "model", "模型配置", "测试模型注册；LLM Mock 或测试 Key", "", "AI 可调用", ""],
    ["D3", "frontend", "系统与仓库", "系统/仓库表单页联调", "", "页面可用", ""],
    ["D4", "scanner", "代码拉取与扫描", "全量 git clone / snapshot", "Git 凭据", "代码已拉取", ""],
    ["D4", "parser", "Java 静态解析", "全量 JavaParser AST 落库", "", "解析完成", ""],
    ["D4", "callchain", "调用链持久化", "ci_method_call 全量写入", "", "调用链就绪", ""],
    ["D4", "entrypoint", "入口识别与复核", "入口自动识别落表（W1 关闭人工断点）", "requireEntrypointReview=false", "", ""],
    ["D4", "hierarchy", "模块层级与复核", "AI 模块层级提炼（W1 关闭层级断点）", "requireHierarchyReview=false", "", ""],
    ["D4", "ai", "AI 归纳调用", "模块层级 + 功能文档 AI 调用", "model / prompt", "", ""],
    ["D4", "task", "任务状态机与调度", "手动下发 INITIAL 全量任务；详情/日志/进度", "", "任务可创建", ""],
    ["D4", "frontend", "任务列表 / 详情 / 手动下发", "下发全量任务；查看执行进度与日志", "", "任务页可用", ""],
    ["D5", "draft", "草稿工作区", "全量任务产出 Markdown 草稿；只读浏览", "", "≥1 篇草稿可读", ""],
    ["D5", "task", "任务状态机与调度", "INITIAL 全量主链路端到端冒烟", "scanner→draft", "全量流程跑通", ""],
    ["D5", "pipeline", "前端部署流水线", "前端静态资源部署测试环境", "", "前端可访问", ""],
    ["D5", "pipeline", "后端部署流水线", "测试环境部署回归；流水线再跑 1 次", "", "CI 跑通 2 次", ""],
    ["D5", "frontend", "生成知识复核", "草稿只读页冒烟（Monaco 预览即可）", "", "W1 演示就绪", ""],
]

W1_EXIT = (
    "W1 退出标准：测试环境前后端部署成功 · Apollo TEST + Bettle DDL · 统一缓存连通 · "
    "PACAS 可登录 · 单仓库 Git 拉代码 · INITIAL 全量任务跑通 · 至少 1 篇草稿可读 · 流水线各跑通"
)

# W2：按分工总览补齐剩余模块/页面 + 生产环境
W2_ROWS = [
    ["D6", "auth", "UAMP 权限对接", "菜单/按钮/API 权限全量", "", "", ""],
    ["D6", "entrypoint", "入口识别与复核", "扫描入口复核页 + 断点开启", "", "", ""],
    ["D6", "hierarchy", "模块层级与复核", "知识模块复核页 + 断点开启", "", "", ""],
    ["D6", "frontend", "扫描入口复核 / 知识模块复核", "双重复核列表与详情页", "", "", ""],
    ["D7", "draft", "草稿工作区", "编辑锁 + 自动保存 + 确认（统一缓存）", "", "", ""],
    ["D7", "draft", "草稿工作区", "Monaco 编辑 + 模块目录树", "", "", ""],
    ["D7", "frontend", "生成知识复核", "草稿工作区全功能", "", "", ""],
    ["D8", "scanner", "代码拉取与扫描", "增量 diff + 增量影响分析", "", "", ""],
    ["D8", "callchain", "调用链持久化", "增量删重建调用链", "", "", ""],
    ["D8", "ai", "AI 归纳调用", "增量模块/文档重跑", "", "", ""],
    ["D8", "knowledge", "知识查看与纠错", "扫描入口 / 模块层级 / 知识文档三页", "", "", ""],
    ["D8", "frontend", "知识查询三页", "入口 / 层级 / 文档 + 纠错抽屉", "", "", ""],
    ["D9", "push", "知识推送与版本", "NAS/Git 发布；版本列表", "", "", ""],
    ["D9", "frontend", "知识推送", "推送确认 + ZIP 导出页", "", "", ""],
    ["D9", "dashboard", "运营看板", "任务概览 / AI 用量 / 流水线分析 / 覆盖率", "", "", ""],
    ["D9", "token", "Token 审计", "Token 审计页 + API", "", "", ""],
    ["D9", "log", "操作日志", "操作日志检索页", "", "", ""],
    ["D9", "frontend", "仪表盘 / 看板", "4 看板页 + Token 审计 + 操作日志", "", "", ""],
    ["D10", "integration", "神兵 Webhook", "真实 Webhook 联调（或生产 Mock 升级）", "task 入队 API", "", ""],
    ["D10", "integration", "CLI 知识聚合 API", "CLI 下载完善", "storage NAS 路径", "", ""],
    ["D10", "integration", "知识打通 ai_workspace", "定时同步验证", "联系人：杨东", "", ""],
    ["D10", "task", "神兵任务入队 API", "与神兵生产联调", "", "", ""],
    ["D10", "quotacontrol", "流量控制", "配额 / 并发配置页", "", "", ""],
    ["D10", "scanwindow", "扫描编排", "任务编排热力图 + 神兵触发", "", "", ""],
    ["D10", "businessknowledge", "业务知识维护", "业务知识 CRUD", "", "", ""],
    ["D10", "frontend", "基础配置剩余页", "模型 / 提示词 / 权限 / 流量 / 编排 / 队列", "", "", ""],
    ["D11", "infra", "Apollo 全量配置", "Apollo PRO 发布", "", "生产配置就绪", ""],
    ["D11", "infra", "统一缓存对接", "生产统一缓存验证", "", "", ""],
    ["D11", "infra", "Bettle + DB 初始化", "生产 DDL 执行", "", "", ""],
    ["D11", "pipeline", "后端部署流水线", "生产环境部署 Job", "", "生产后端可访问", ""],
    ["D11", "pipeline", "前端部署流水线", "生产静态资源部署", "", "生产前端可访问", ""],
    ["D12", "全模块", "UAT 与回归", "INITIAL + INCREMENTAL + 双重复核 + 推送 + 纠错全路径", "", "UAT 通过", ""],
]

W2_EXIT = (
    "W2 退出标准：分工总览剩余模块/页面补齐 · 增量扫描 · 知识三页+纠错+推送 · "
    "神兵/CLI/ai_workspace · 生产环境部署 · UAT 通过"
)


def style_header_row(ws, row, ncol):
    for c in range(1, ncol + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER


def write_section(ws, start_row, title, data_rows):
    ncol = len(HEADERS)
    ws.merge_cells(start_row=start_row, start_column=1, end_row=start_row, end_column=ncol)
    cell = ws.cell(row=start_row, column=1, value=title)
    cell.fill = SECTION_FILL
    cell.font = SECTION_FONT
    for c in range(1, ncol + 1):
        ws.cell(row=start_row, column=c).border = BORDER

    r = start_row + 1
    for i, h in enumerate(HEADERS, 1):
        ws.cell(row=r, column=i, value=h)
    style_header_row(ws, r, ncol)
    r += 1

    for row_data in data_rows:
        for i, val in enumerate(row_data, 1):
            c = ws.cell(row=r, column=i, value=val)
            c.alignment = WRAP
            c.border = BORDER
        r += 1
    return r


def build_schedule_sheet(ws):
    ws.delete_rows(1, ws.max_row)
    ws["A1"] = "CodeInsight 公司框架迁移 — 两周排期"
    ws.merge_cells("A1:G1")
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = "第1周：测试环境部署 + 全量扫描主链路｜第2周：分工总览剩余项 + 生产上线"
    ws.merge_cells("A2:G2")
    ws["A2"].font = Font(size=10, color="666666")

    row = 4
    row = write_section(ws, row, "第 1 周 — 测试环境部署 + 全量扫描流程打通", W1_ROWS)
    row += 1
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=len(HEADERS))
    ws.cell(row=row, column=1, value=W1_EXIT).font = Font(bold=True, color="C65911")
    ws.cell(row=row, column=1).alignment = WRAP
    row += 2

    row = write_section(ws, row, "第 2 周 — 全功能补齐 + 生产部署", W2_ROWS)
    row += 1
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=len(HEADERS))
    ws.cell(row=row, column=1, value=W2_EXIT).font = Font(bold=True, color="006100")
    ws.cell(row=row, column=1).alignment = WRAP

    for i, w in enumerate(COL_WIDTHS, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A5"


def main():
    for path in (XLSX, XLSX_CN):
        wb = load_workbook(path)
        if "排期" in wb.sheetnames:
            ws = wb["排期"]
        else:
            ws = wb.create_sheet("排期")
        build_schedule_sheet(ws)
        wb.save(path)
        print(f"Updated schedule in: {path}")


if __name__ == "__main__":
    main()
